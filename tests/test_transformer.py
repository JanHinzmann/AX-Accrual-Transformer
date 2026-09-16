from __future__ import annotations

import csv
import io
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from transformer import AX_HEADERS, WorkbookValidationError, process_accrual_workbook


SOURCE_HEADERS = [
    "Company ID",
    "Cost center ID",
    "Ledger account ID",
    "Project ID",
    "Vendor ID",
    "Vendor name",
    "Posting description",
    "Time",
    "Value",
]


def source_row(company_id=120, value=100, **overrides):
    values = {
        "Company ID": company_id,
        "Cost center ID": 501300,
        "Ledger account ID": 681200000,
        "Project ID": "PR-01",
        "Vendor ID": "V-007",
        "Vendor name": "Example Vendor",
        "Posting description": "Consulting fee",
        "Time": 8,
        "Value": value,
    }
    values.update(overrides)
    return values


def workbook_bytes(rows, *, headers=SOURCE_HEADERS, header_row=1):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DATA"
    for row_index in range(1, header_row):
        worksheet.cell(row_index, 1, "Accrual upload")
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column_index, header)
    for row in rows:
        worksheet.append([row.get(header) if isinstance(row, dict) else None for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def write_source_sheet(worksheet, rows, *, headers=SOURCE_HEADERS, header_row=1):
    for row_index in range(1, header_row):
        worksheet.cell(row_index, 1, "Accrual upload")
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column_index, header)
    for row in rows:
        worksheet.append([row.get(header) if isinstance(row, dict) else None for header in headers])


def save_workbook_bytes(workbook):
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def write_mapping(path: Path, rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DATA"
    worksheet.append(["Company ID", "ERP System"])
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    return path


def write_blank_template(path: Path):
    workbook = Workbook()
    workbook.active.title = "Tabelle1"
    workbook.active["A1"].fill = PatternFill("solid", fgColor="ED0334")
    workbook.save(path)
    return path


def run_process(
    tmp_path: Path,
    rows,
    mappings,
    *,
    template=None,
    name="source.xlsx",
    header_row=1,
    headers=SOURCE_HEADERS,
    delimiter=";",
):
    mapping_path = write_mapping(tmp_path / "mapping.xlsx", mappings)
    template_path = template or write_blank_template(tmp_path / "template.xlsx")
    return process_accrual_workbook(
        workbook_bytes(rows, headers=headers, header_row=header_row),
        original_filename=name,
        company_mapping_path=mapping_path,
        ax_template_path=template_path,
        processing_date=date(2026, 9, 16),
        csv_delimiter=delimiter,
    )


def parse_generated(result, company_id, *, delimiter=";"):
    generated = next(item for item in result.files if item.company_id == str(company_id))
    text = generated.content.decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    headers = parsed[0]
    rows = [dict(zip(headers, row)) for row in parsed[1:]]
    return generated, headers, rows, text


def run_process_bytes(
    tmp_path: Path,
    uploaded,
    mappings,
    *,
    template=None,
    name="source.xlsx",
    delimiter=";",
):
    mapping_path = write_mapping(tmp_path / "mapping.xlsx", mappings)
    template_path = template or write_blank_template(tmp_path / "template.xlsx")
    return process_accrual_workbook(
        uploaded,
        original_filename=name,
        company_mapping_path=mapping_path,
        ax_template_path=template_path,
        processing_date=date(2026, 9, 16),
        csv_delimiter=delimiter,
    )


def test_positive_and_negative_values_populate_correct_sides(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(value=125.5), source_row(value=-48.25)],
        [(120, "AX")],
    )
    _, _, rows, _ = parse_generated(result, 120)

    assert rows[0]["Soll"] == "125,50"
    assert rows[0]["Haben"] == ""
    assert rows[0]["Gegenkonto"] == "307030000"
    assert rows[1]["Soll"] == ""
    assert rows[1]["Haben"] == "48,25"
    assert rows[1]["Gegenkonto"] == "130010010"


def test_posting_text_numeric_parsing_and_identifiers(tmp_path):
    row = source_row(
        value="1.234,56",
        **{
            "Time": " 08 ",
            "Vendor name": " Vendor GmbH ",
            "Posting description": " Service ",
            "Vendor ID": "0007",
        },
    )
    result = run_process(tmp_path, [row], [(120, "AX")])
    _, _, rows, _ = parse_generated(result, 120)

    assert rows[0]["Buchungstext"] == "08.2026 RUECK Vendor GmbH Service , 0007"
    assert rows[0]["Soll"] == "1234,56"
    assert rows[0]["Konto"] == "681200000"


@pytest.mark.parametrize("time_value", [None, "", "   ", float("nan")])
def test_blank_time_builds_posting_description_and_vendor_id(tmp_path, time_value):
    row = source_row(
        **{
            "Time": time_value,
            "Posting description": "Audit fee",
            "Vendor ID": "0007",
        }
    )
    result = run_process(tmp_path, [row], [(120, "AX")])
    _, _, rows, _ = parse_generated(result, 120)

    assert rows[0]["Buchungstext"] == "Audit fee , 0007"


@pytest.mark.parametrize(
    ("time_value", "expected_prefix"),
    [
        (1, "01.2026 RUECK"),
        ("7", "07.2026 RUECK"),
        (9.0, "09.2026 RUECK"),
        ("09", "09.2026 RUECK"),
        (12, "12.2026 RUECK"),
    ],
)
def test_posting_text_formats_time_without_double_padding(tmp_path, time_value, expected_prefix):
    result = run_process(tmp_path, [source_row(**{"Time": time_value})], [(120, "AX")])
    _, _, rows, _ = parse_generated(result, 120)

    assert rows[0]["Buchungstext"].startswith(expected_prefix)


def test_posting_text_omits_missing_optional_values(tmp_path):
    row = source_row(
        **{
            "Time": 4,
            "Vendor name": None,
            "Posting description": "Audit fee",
            "Vendor ID": None,
        }
    )
    result = run_process(tmp_path, [row], [(120, "AX")])
    _, _, rows, _ = parse_generated(result, 120)

    posting_text = rows[0]["Buchungstext"]
    assert posting_text.startswith("04.2026 RUECK")
    assert "nan" not in posting_text.casefold()
    assert "none" not in posting_text.casefold()


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (125, "125,00"),
        (125.5, "125,50"),
        ("125.554", "125,55"),
        ("125.555", "125,56"),
    ],
)
def test_csv_amounts_use_decimal_comma_two_places_and_half_up_rounding(
    tmp_path, value, expected
):
    result = run_process(tmp_path, [source_row(value=value)], [(120, "AX")])
    _, _, rows, _ = parse_generated(result, 120)

    assert rows[0]["Soll"] == expected
    assert rows[0]["Haben"] == ""


def test_split_files_are_csv_and_index_restarts_for_each_company(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(120, 10), source_row(120, 20), source_row(140, 30)],
        [(120, "AX"), (140, "AX")],
    )
    assert {item.company_id for item in result.files} == {"120", "140"}
    assert all(item.filename.endswith(".csv") for item in result.files)
    for company_id in (120, 140):
        _, _, rows, _ = parse_generated(result, company_id)
        assert rows[0]["autom. Vergabe von Belegnummern"] == "1"


def test_mixed_ax_and_d365_generates_ax_and_skips_d365(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(120, 10), source_row(100, 20)],
        [(120, "AX"), (100, "D365")],
    )
    assert [item.company_id for item in result.files] == ["120"]
    outcomes = {item.company_id: item for item in result.outcomes}
    assert outcomes["120"].status == "Generated"
    assert outcomes["100"].status == "Skipped"
    assert "not implemented" in outcomes["100"].message


def test_blank_rows_do_not_create_csv_records(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(120, 10), {}, source_row(120, 20)],
        [(120, "AX")],
    )
    _, _, rows, _ = parse_generated(result, 120)

    assert result.source_row_count == 2
    assert [row["autom. Vergabe von Belegnummern"] for row in rows] == ["1", "2"]


def test_unmapped_and_conflicting_company_ids_fail_independently(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(120, 10), source_row(140, 20)],
        [(140, "AX"), (140, "D365")],
    )
    outcomes = {item.company_id: item for item in result.outcomes}

    assert result.files == []
    assert outcomes["120"].status == "Failed"
    assert "not found" in outcomes["120"].message
    assert outcomes["140"].status == "Failed"
    assert "conflicting" in outcomes["140"].message


def test_invalid_and_zero_values_block_only_the_affected_company(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(120, 0), source_row(120, "not a number"), source_row(140, 25)],
        [(120, "AX"), (140, "AX")],
    )
    outcomes = {item.company_id: item for item in result.outcomes}

    assert [item.company_id for item in result.files] == ["140"]
    assert outcomes["120"].status == "Failed"
    assert outcomes["140"].status == "Generated"
    assert {issue.source_row for issue in result.issues if issue.company_id == "120"} == {2, 3}


def test_reversal_date_is_written_in_german_date_format(tmp_path):
    result = run_process(tmp_path, [source_row()], [(120, "AX")])
    _, _, rows, _ = parse_generated(result, 120)

    assert rows[0]["Rückbuchungsdatum"] == "01.09.2026"


def test_default_csv_has_utf8_bom_crlf_and_exact_headers(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(**{"Vendor name": "Österreichische Apotheker"})],
        [(120, "AX")],
    )
    generated, headers, rows, text = parse_generated(result, 120)

    assert generated.content.startswith(b"\xef\xbb\xbf")
    assert headers == list(AX_HEADERS)
    assert len(headers) == 19
    assert "\r\n" in text
    assert rows[0]["Buchungstext"].startswith("08.2026 RUECK Österreichische Apotheker")


def test_populated_template_controls_csv_header_order(tmp_path):
    template = tmp_path / "populated-template.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "AX Import"
    worksheet["A1"] = "Template title"
    reversed_headers = list(reversed(AX_HEADERS))
    for column, header in enumerate(reversed_headers, start=1):
        worksheet.cell(2, column, header)
    workbook.save(template)

    result = run_process(tmp_path, [source_row()], [(120, "AX")], template=template)
    _, headers, rows, _ = parse_generated(result, 120)

    assert headers == reversed_headers
    assert rows[0]["Kostenart"] == "Sachkonto"
    assert rows[0]["autom. Vergabe von Belegnummern"] == "1"


def test_headers_can_be_detected_below_first_row(tmp_path):
    result = run_process(tmp_path, [source_row()], [(120, "AX")], header_row=3)
    assert result.source_sheet == "DATA"
    assert result.source_header_row == 3
    assert len(result.files) == 1


def test_extra_tabs_are_ignored_and_trimmed_data_name_is_detected(tmp_path):
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Cover"
    cover.append(["Company ID", "Company ID", "Notes"])
    data = workbook.create_sheet(" DATA ")
    write_source_sheet(data, [source_row(**{"Time": 3})], header_row=3)
    notes = workbook.create_sheet("Notes")
    notes.append(["Time", "Unrelated content"])

    result = run_process_bytes(
        tmp_path,
        save_workbook_bytes(workbook),
        [(120, "AX")],
    )
    _, _, rows, _ = parse_generated(result, 120)

    assert result.source_sheet == " DATA "
    assert result.source_header_row == 3
    assert rows[0]["Buchungstext"].startswith("03.2026 RUECK")


def test_valid_data_sheet_is_preferred_over_an_earlier_valid_sheet(tmp_path):
    workbook = Workbook()
    archive = workbook.active
    archive.title = "Archive"
    write_source_sheet(archive, [source_row(**{"Time": 11})])
    data = workbook.create_sheet("data")
    write_source_sheet(data, [source_row(**{"Time": 3})])

    result = run_process_bytes(
        tmp_path,
        save_workbook_bytes(workbook),
        [(120, "AX")],
    )
    _, _, rows, _ = parse_generated(result, 120)

    assert result.source_sheet == "data"
    assert result.source_row_count == 1
    assert rows[0]["Buchungstext"].startswith("03.2026 RUECK")


def test_first_valid_sheet_is_used_when_there_is_no_data_sheet(tmp_path):
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Cover"
    cover.append(["Instructions only"])
    current = workbook.create_sheet("Current accruals")
    write_source_sheet(current, [source_row(**{"Time": 6})])
    archive = workbook.create_sheet("Archive")
    write_source_sheet(archive, [source_row(**{"Time": 12})])

    result = run_process_bytes(
        tmp_path,
        save_workbook_bytes(workbook),
        [(120, "AX")],
    )
    _, _, rows, _ = parse_generated(result, 120)

    assert result.source_sheet == "Current accruals"
    assert rows[0]["Buchungstext"].startswith("06.2026 RUECK")


def test_extra_columns_and_old_month_are_ignored(tmp_path):
    headers = [
        "Extra Before",
        "Time",
        "Company ID",
        "Cost center ID",
        "Extra Between",
        "Ledger account ID",
        "Project ID",
        "Vendor ID",
        "Vendor name",
        "Posting description",
        "Month",
        "Value",
        "Extra After",
    ]
    row = source_row(**{"Time": 3})
    row.update(
        {
            "Month": 12,
            "Extra Before": "ignore me",
            "Extra Between": "ignore me too",
            "Extra After": 999,
        }
    )
    irrelevant_only_row = {
        "Extra Before": "not a source row",
        "Month": 10,
        "Extra After": "still irrelevant",
    }

    result = run_process(
        tmp_path,
        [row, irrelevant_only_row],
        [(120, "AX")],
        headers=headers,
    )
    _, generated_headers, rows, _ = parse_generated(result, 120)

    assert result.source_row_count == 1
    assert generated_headers == list(AX_HEADERS)
    assert rows[0]["Buchungstext"].startswith("03.2026 RUECK")
    assert not {"Extra Before", "Extra Between", "Month", "Extra After"} & set(
        generated_headers
    )


def test_month_only_source_fails_with_missing_time_error(tmp_path):
    headers = ["Month" if header == "Time" else header for header in SOURCE_HEADERS]
    row = source_row()
    row["Month"] = row.pop("Time")

    with pytest.raises(WorkbookValidationError) as exc_info:
        run_process(tmp_path, [row], [(120, "AX")], headers=headers)

    message = str(exc_info.value)
    assert "Missing required columns: Time" in message
    assert "Expected:" in message


def test_duplicate_required_header_on_selected_source_row_fails(tmp_path):
    headers = [*SOURCE_HEADERS, "Time"]

    with pytest.raises(WorkbookValidationError) as exc_info:
        run_process(tmp_path, [source_row()], [(120, "AX")], headers=headers)

    message = str(exc_info.value)
    assert "header 'Time' more than once" in message
    assert "sheet 'DATA'" in message


def test_missing_required_source_column_has_clear_error(tmp_path):
    mapping = write_mapping(tmp_path / "mapping.xlsx", [(120, "AX")])
    template = write_blank_template(tmp_path / "template.xlsx")
    headers = [header for header in SOURCE_HEADERS if header != "Value"]

    with pytest.raises(WorkbookValidationError) as exc_info:
        process_accrual_workbook(
            workbook_bytes([source_row()], headers=headers),
            original_filename="source.xlsx",
            company_mapping_path=mapping,
            ax_template_path=template,
            processing_date=date(2026, 9, 16),
        )

    assert "Expected:" in str(exc_info.value)
    assert "Value" in str(exc_info.value)
    assert "Detected:" in str(exc_info.value)


def test_zip_contains_only_generated_csv_files(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(120, 10), source_row(140, 20)],
        [(120, "AX"), (140, "AX")],
        name="ACCRUAL DATA.xlsx",
    )
    with zipfile.ZipFile(io.BytesIO(result.as_zip())) as archive:
        assert sorted(archive.namelist()) == sorted(item.filename for item in result.files)
        assert all(name.endswith(".csv") for name in archive.namelist())


def test_comma_delimiter_quotes_commas_and_round_trips(tmp_path):
    result = run_process(
        tmp_path,
        [source_row(**{"Vendor name": "Vendor, Inc."})],
        [(120, "AX")],
        delimiter=",",
    )
    _, headers, rows, text = parse_generated(result, 120, delimiter=",")

    assert headers == list(AX_HEADERS)
    assert rows[0]["Buchungstext"] == "08.2026 RUECK Vendor, Inc. Consulting fee , V-007"
    assert rows[0]["Soll"] == "100,00"
    assert '"08.2026 RUECK Vendor, Inc. Consulting fee , V-007"' in text
    assert '"100,00"' in text
    parsed = list(csv.reader(io.StringIO(text, newline=""), delimiter=","))
    assert all(len(record) == len(headers) for record in parsed)
