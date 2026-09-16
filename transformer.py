from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO, Iterable
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


REQUIRED_SOURCE_COLUMNS = (
    "Company ID",
    "Ledger account ID",
    "Cost center ID",
    "Value",
    "Month",
    "Vendor name",
    "Posting description",
    "Vendor ID",
    "Project ID",
)

AX_HEADERS = (
    "autom. Vergabe von Belegnummern",
    "Kostenart",
    "Konto",
    "Rechnung",
    "MWSt Gruppe",
    "Artikel MWSt Gr.",
    "KST",
    "Soll",
    "Haben",
    "Buchungstext",
    "Gegenkontenart",
    "Gegenkonto",
    "PaymMode",
    "Payment",
    "Document Num",
    "Document Date",
    "Projekt",
    "Rückbuchung (ja/nein)",
    "Rückbuchungsdatum",
)

MAPPING_COLUMNS = ("Company ID", "ERP System")
BERLIN = ZoneInfo("Europe/Berlin")


class WorkbookValidationError(ValueError):
    """Raised when an uploaded or reference workbook is structurally invalid."""


@dataclass(frozen=True)
class RowIssue:
    company_id: str
    source_row: int
    field: str
    message: str


@dataclass(frozen=True)
class CompanyOutcome:
    company_id: str
    erp_system: str
    source_row_count: int
    status: str
    message: str


@dataclass(frozen=True)
class GeneratedFile:
    company_id: str
    erp_system: str
    filename: str
    content: bytes


@dataclass
class ProcessingResult:
    source_sheet: str
    source_header_row: int
    source_row_count: int
    outcomes: list[CompanyOutcome] = field(default_factory=list)
    issues: list[RowIssue] = field(default_factory=list)
    files: list[GeneratedFile] = field(default_factory=list)

    def as_zip(self) -> bytes:
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for generated in sorted(self.files, key=lambda item: item.filename.casefold()):
                archive.writestr(generated.filename, generated.content)
        return output.getvalue()


@dataclass(frozen=True)
class _SourceRow:
    row_number: int
    company_id: str
    values: dict[str, object]
    number_formats: dict[str, str]


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split()).casefold()


def _is_blank(value: object) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    if isinstance(value, float):
        return not math.isfinite(value)
    if isinstance(value, Decimal):
        return not value.is_finite()
    return False


def _clean_text(value: object) -> str:
    if _is_blank(value):
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        return format(value, "f").rstrip("0").rstrip(".") if value % 1 else str(int(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    text = str(value).strip()
    if text.casefold() in {"nan", "none"}:
        return ""
    return text


def _copy_text(value: object) -> str:
    """Copy source text without changing wording or surrounding whitespace."""
    if _is_blank(value):
        return ""
    if isinstance(value, str):
        return "" if value.strip().casefold() in {"nan", "none"} else value
    return _clean_text(value)


def _identifier_text(value: object, number_format: str = "") -> str:
    """Return a stable identifier without artificial .0 or lost displayed zeroes."""
    text = _clean_text(value)
    if not text:
        return ""

    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        simple_format = number_format.strip().split(";")[0]
        if re.fullmatch(r"0+", simple_format):
            integer = Decimal(str(value))
            if integer == integer.to_integral_value():
                return str(int(integer)).zfill(len(simple_format))

    match = re.fullmatch(r"([+-]?\d+)\.0+", text)
    return match.group(1) if match else text


def _parse_decimal(value: object) -> Decimal | None:
    if _is_blank(value) or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value)) if math.isfinite(value) else None

    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace("'", "")
    negative_parentheses = text.startswith("(") and text.endswith(")")
    if negative_parentheses:
        text = text[1:-1]
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        parsed = Decimal(text)
    except InvalidOperation:
        return None
    if not parsed.is_finite():
        return None
    return -parsed if negative_parentheses else parsed


def _locate_header(
    workbook,
    required_columns: Iterable[str],
    *,
    workbook_label: str,
) -> tuple[object, int, dict[str, int]]:
    required = tuple(required_columns)
    normalized_required = {_normalize_header(name): name for name in required}
    detected_samples: list[str] = []

    for worksheet in workbook.worksheets:
        max_scan_row = min(max(worksheet.max_row, 1), 50)
        max_scan_col = min(max(worksheet.max_column, 1), 100)
        for row_index in range(1, max_scan_row + 1):
            found: dict[str, int] = {}
            displayed: list[str] = []
            for column_index in range(1, max_scan_col + 1):
                value = worksheet.cell(row_index, column_index).value
                normalized = _normalize_header(value)
                if normalized:
                    displayed.append(_clean_text(value))
                    if normalized in normalized_required:
                        if normalized in found:
                            raise WorkbookValidationError(
                                f"{workbook_label} contains the header "
                                f"'{normalized_required[normalized]}' more than once on "
                                f"sheet '{worksheet.title}', row {row_index}."
                            )
                        found[normalized] = column_index
            if displayed and len(detected_samples) < 5:
                detected_samples.append(f"{worksheet.title}!{row_index}: {', '.join(displayed)}")
            if all(normalized in found for normalized in normalized_required):
                return (
                    worksheet,
                    row_index,
                    {name: found[normalized] for normalized, name in normalized_required.items()},
                )

    expected = ", ".join(required)
    detected = "; ".join(detected_samples) if detected_samples else "no populated header candidates"
    raise WorkbookValidationError(
        f"Could not find all required columns in {workbook_label}. "
        f"Expected: {expected}. Detected: {detected}."
    )


def _load_company_mapping(mapping_path: Path) -> dict[str, set[str]]:
    try:
        workbook = load_workbook(mapping_path, data_only=True, read_only=False)
    except Exception as exc:
        raise WorkbookValidationError(f"Could not read COMPANY MAPPING.xlsx: {exc}") from exc

    worksheet, header_row, columns = _locate_header(
        workbook, MAPPING_COLUMNS, workbook_label="COMPANY MAPPING.xlsx"
    )
    mappings: dict[str, set[str]] = {}
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        company_cell = worksheet.cell(row_index, columns["Company ID"])
        erp_cell = worksheet.cell(row_index, columns["ERP System"])
        company_id = _identifier_text(company_cell.value, company_cell.number_format)
        erp = _clean_text(erp_cell.value).upper()
        if not company_id and not erp:
            continue
        if not company_id or not erp:
            continue
        mappings.setdefault(company_id, set()).add(erp)
    return mappings


def _read_source(uploaded: bytes | BinaryIO) -> tuple[str, int, list[_SourceRow], list[RowIssue]]:
    payload = io.BytesIO(uploaded) if isinstance(uploaded, bytes) else uploaded
    try:
        workbook = load_workbook(payload, data_only=True, read_only=False)
    except Exception as exc:
        raise WorkbookValidationError(
            "The uploaded file could not be read as an .xlsx workbook. "
            f"Technical detail: {exc}"
        ) from exc

    worksheet, header_row, columns = _locate_header(
        workbook, REQUIRED_SOURCE_COLUMNS, workbook_label="the uploaded accrual workbook"
    )
    rows: list[_SourceRow] = []
    issues: list[RowIssue] = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        cells = {name: worksheet.cell(row_index, column) for name, column in columns.items()}
        if all(_is_blank(cell.value) for cell in cells.values()):
            continue
        company_cell = cells["Company ID"]
        company_id = _identifier_text(company_cell.value, company_cell.number_format)
        if not company_id:
            issues.append(
                RowIssue("(blank)", row_index, "Company ID", "Company ID is missing; the row was skipped.")
            )
            continue
        rows.append(
            _SourceRow(
                row_number=row_index,
                company_id=company_id,
                values={name: cell.value for name, cell in cells.items()},
                number_formats={name: cell.number_format for name, cell in cells.items()},
            )
        )
    return worksheet.title, header_row, rows, issues


def _build_posting_text(row: _SourceRow) -> str:
    month = _identifier_text(row.values["Month"], row.number_formats["Month"])
    if not month:
        return _copy_text(row.values["Posting description"])
    if re.fullmatch(r"[1-9]", month):
        month = month.zfill(2)

    vendor_name = _clean_text(row.values["Vendor name"])
    posting_description = _clean_text(row.values["Posting description"])
    vendor_id = _identifier_text(row.values["Vendor ID"], row.number_formats["Vendor ID"])
    return f"{month}.2026 RUECK {vendor_name} {posting_description} , {vendor_id}".strip()


def _validate_ax_rows(rows: list[_SourceRow]) -> tuple[dict[int, Decimal], list[RowIssue]]:
    parsed_values: dict[int, Decimal] = {}
    issues: list[RowIssue] = []
    for row in rows:
        ledger = _identifier_text(
            row.values["Ledger account ID"], row.number_formats["Ledger account ID"]
        )
        cost_center = _identifier_text(
            row.values["Cost center ID"], row.number_formats["Cost center ID"]
        )
        if not ledger:
            issues.append(
                RowIssue(row.company_id, row.row_number, "Ledger account ID", "Ledger account ID is missing.")
            )
        if not cost_center:
            issues.append(
                RowIssue(row.company_id, row.row_number, "Cost center ID", "Cost center ID is missing.")
            )

        value = _parse_decimal(row.values["Value"])
        if value is None:
            issues.append(
                RowIssue(row.company_id, row.row_number, "Value", "Value is missing or is not numeric.")
            )
        elif value == 0:
            issues.append(
                RowIssue(row.company_id, row.row_number, "Value", "Value is zero; no balanced AX row can be created.")
            )
        else:
            parsed_values[row.row_number] = value

        if not _build_posting_text(row).strip():
            issues.append(
                RowIssue(
                    row.company_id,
                    row.row_number,
                    "Buchungstext",
                    "Buchungstext is empty; Posting description is required when Month is empty.",
                )
            )
    return parsed_values, issues


def _resolve_ax_header_order(template_path: Path) -> list[str]:
    try:
        workbook = load_workbook(template_path, data_only=False, read_only=False)
    except Exception as exc:
        raise WorkbookValidationError(f"Could not read AX TEMPLATE.xlsx: {exc}") from exc

    try:
        worksheet, header_row, columns = _locate_header(
            workbook, AX_HEADERS, workbook_label="AX TEMPLATE.xlsx"
        )
        del worksheet, header_row
        return sorted(AX_HEADERS, key=lambda header: columns[header])
    except WorkbookValidationError:
        populated = any(
            not _is_blank(cell.value)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        )
        if populated:
            raise
        return list(AX_HEADERS)


def _format_csv_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return str(value)


def _write_ax_csv(
    rows: list[_SourceRow],
    parsed_values: dict[int, Decimal],
    template_path: Path,
    reversal_date: date,
    delimiter: str,
) -> bytes:
    if delimiter not in {";", ","}:
        raise ValueError("CSV delimiter must be ';' or ','.")
    headers = _resolve_ax_header_order(template_path)
    output = io.StringIO(newline="")
    writer = csv.writer(
        output,
        delimiter=delimiter,
        quotechar='"',
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\r\n",
    )
    writer.writerow(headers)

    for offset, source_row in enumerate(rows):
        value = parsed_values[source_row.row_number]
        posting_text = _build_posting_text(source_row)
        debit = value if value > 0 else None
        credit = abs(value) if value < 0 else None
        row_values = {
            "autom. Vergabe von Belegnummern": offset + 1,
            "Kostenart": "Sachkonto",
            "Konto": _identifier_text(
                source_row.values["Ledger account ID"],
                source_row.number_formats["Ledger account ID"],
            ),
            "Rechnung": None,
            "MWSt Gruppe": None,
            "Artikel MWSt Gr.": None,
            "KST": _identifier_text(
                source_row.values["Cost center ID"], source_row.number_formats["Cost center ID"]
            ),
            "Soll": debit,
            "Haben": credit,
            "Buchungstext": posting_text,
            "Gegenkontenart": "Sachkonto",
            "Gegenkonto": "130010010" if credit is not None else "307030000",
            "PaymMode": None,
            "Payment": None,
            "Document Num": None,
            "Document Date": None,
            "Projekt": _identifier_text(
                source_row.values["Project ID"], source_row.number_formats["Project ID"]
            ),
            "Rückbuchung (ja/nein)": "ja",
            "Rückbuchungsdatum": reversal_date,
        }
        writer.writerow([_format_csv_value(row_values[header]) for header in headers])

    return output.getvalue().encode("utf-8-sig")


def _safe_filename_part(value: str, *, default: str) -> str:
    sanitized = re.sub(r"[<>:\"/\\|?*\x00-\x1F]+", "_", value.strip())
    sanitized = re.sub(r"\s+", "_", sanitized).strip(" ._")
    return (sanitized or default)[:100]


def process_accrual_workbook(
    uploaded: bytes | BinaryIO,
    *,
    original_filename: str,
    company_mapping_path: str | Path,
    ax_template_path: str | Path,
    processing_date: date | datetime | None = None,
    csv_delimiter: str = ";",
) -> ProcessingResult:
    """Validate, split, and transform one uploaded accrual workbook."""
    source_sheet, source_header_row, rows, initial_issues = _read_source(uploaded)
    mappings = _load_company_mapping(Path(company_mapping_path))
    current_date = processing_date
    if current_date is None:
        current_date = datetime.now(BERLIN)
    reversal_date = date(current_date.year, current_date.month, 1)

    grouped: dict[str, list[_SourceRow]] = {}
    for row in rows:
        grouped.setdefault(row.company_id, []).append(row)

    result = ProcessingResult(
        source_sheet=source_sheet,
        source_header_row=source_header_row,
        source_row_count=len(rows) + len(initial_issues),
        issues=list(initial_issues),
    )
    if initial_issues:
        result.outcomes.append(
            CompanyOutcome(
                company_id="(blank)",
                erp_system="",
                source_row_count=len(initial_issues),
                status="Failed",
                message="Rows without Company ID were skipped.",
            )
        )

    source_stem = _safe_filename_part(Path(original_filename).stem, default="accrual_data")
    for company_id, company_rows in grouped.items():
        erps = mappings.get(company_id)
        if not erps:
            result.outcomes.append(
                CompanyOutcome(
                    company_id,
                    "",
                    len(company_rows),
                    "Failed",
                    "Company ID was not found in COMPANY MAPPING.xlsx.",
                )
            )
            continue
        if len(erps) > 1:
            result.outcomes.append(
                CompanyOutcome(
                    company_id,
                    ", ".join(sorted(erps)),
                    len(company_rows),
                    "Failed",
                    "Company ID has conflicting ERP mappings.",
                )
            )
            continue

        erp = next(iter(erps))
        if erp == "D365":
            result.outcomes.append(
                CompanyOutcome(
                    company_id,
                    erp,
                    len(company_rows),
                    "Skipped",
                    "D365 transformation is not implemented yet.",
                )
            )
            continue
        if erp != "AX":
            result.outcomes.append(
                CompanyOutcome(
                    company_id,
                    erp,
                    len(company_rows),
                    "Failed",
                    f"ERP system '{erp}' is not supported.",
                )
            )
            continue

        parsed_values, company_issues = _validate_ax_rows(company_rows)
        if company_issues:
            result.issues.extend(company_issues)
            result.outcomes.append(
                CompanyOutcome(
                    company_id,
                    erp,
                    len(company_rows),
                    "Failed",
                    f"{len(company_issues)} row validation issue(s); no workbook was generated.",
                )
            )
            continue

        content = _write_ax_csv(
            company_rows,
            parsed_values,
            Path(ax_template_path),
            reversal_date,
            csv_delimiter,
        )
        filename = f"AX_{_safe_filename_part(company_id, default='company')}_{source_stem}.csv"
        result.files.append(GeneratedFile(company_id, erp, filename, content))
        result.outcomes.append(
            CompanyOutcome(
                company_id,
                erp,
                len(company_rows),
                "Generated",
                filename,
            )
        )

    return result
